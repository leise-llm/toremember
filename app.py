from flask import Flask, render_template, request, send_file
import openai
import os
from docx import Document
from io import BytesIO

app = Flask(__name__)

# Setze den OpenAI API-Schlüssel aus den Umgebungsvariablen
api_key = os.getenv('OPENAI_API_KEY')
if not api_key:
    raise ValueError("OPENAI_API_KEY ist nicht gesetzt. Bitte überprüfe deine Azure-Konfiguration.")
openai.api_key = api_key

# Globale Variablen, um den Zustand zwischen den Anfragen zu verwalten
reminder_text = ""
refined_text = ""
followup_questions = []
final_text = ""

@app.route('/', methods=['GET', 'POST'])
def index():
    global reminder_text, refined_text, followup_questions, final_text

    if request.method == 'POST':
        # Stichworte aus dem Formular auslesen und Eingaben überprüfen
        keywords = [request.form.get(f'keyword{i}', '').strip() for i in range(1, 6)]
        keywords = [kw for kw in keywords if kw]  # Leere Stichworte entfernen

        if not keywords:
            reminder_text = "Bitte geben Sie mindestens ein Stichwort ein."
        else:
            prompt = (
                "Du bist ein freundlicher Zuhörer. Erstelle einen kurzen Erinnerungstext "
                "in Ich-Form, basierend auf folgenden Stichworten: "
                + ", ".join(keywords) +
                ". Nutze persönliche Sprache ('Ich', 'Wir', 'Du') und beschreibe eine Szene."
            )

            try:
                # Anfrage an OpenAI senden
                response = openai.ChatCompletion.create(
                    model="gpt-4",
                    messages=[
                        {"role": "system", "content": "Du bist ein freundlicher Zuhörer."},
                        {"role": "user", "content": prompt}
                    ],
                    max_tokens=300
                )
                reminder_text = response['choices'][0]['message']['content'].strip()
            except Exception as e:
                reminder_text = f"Es gab ein Problem bei der Anfrage an OpenAI: {str(e)}"

        return render_template('index.html', reminder_text=reminder_text, refinement_prompt=True)

    return render_template('index.html')

@app.route('/refine_reminder', methods=['POST'])
def refine_reminder():
    global reminder_text, refined_text

    user_feedback = request.form.get('user_feedback', '').strip()
    if user_feedback:
        prompt = (
            "Verbessere den folgenden Erinnerungstext mit diesen zusätzlichen Informationen: "
            + reminder_text + " Zusätzliche Informationen: " + user_feedback
        )

        try:
            response = openai.ChatCompletion.create(
                model="gpt-4",
                messages=[
                    {"role": "system", "content": "Du bist ein freundlicher Zuhörer."},
                    {"role": "user", "content": prompt}
                ],
                max_tokens=300
            )
            refined_text = response['choices'][0]['message']['content'].strip()
        except Exception as e:
            refined_text = f"Es gab ein Problem bei der Anfrage an OpenAI: {str(e)}"

    return render_template('index.html', reminder_text=reminder_text, refined_text=refined_text, followup_questions=["Was war das Wetter?", "Wer war dabei?", "Was hast du gefühlt?"])

@app.route('/finalize_reminder', methods=['POST'])
def finalize_reminder():
    global refined_text, final_text

    answers = [request.form.get(f'answer{i}', '').strip() for i in range(1, 4)]
    answers = [ans for ans in answers if ans]  # Leere Antworten entfernen

    if answers:
        prompt = (
            "Integriere diese zusätzlichen Informationen in die erweiterte Erinnerung: "
            + refined_text + " Zusätzliche Informationen: " + ", ".join(answers)
        )

        try:
            response = openai.ChatCompletion.create(
                model="gpt-4",
                messages=[
                    {"role": "system", "content": "Du bist ein freundlicher Zuhörer."},
                    {"role": "user", "content": prompt}
                ],
                max_tokens=300
            )
            final_text = response['choices'][0]['message']['content'].strip()
        except Exception as e:
            final_text = f"Es gab ein Problem bei der Anfrage an OpenAI: {str(e)}"

    return render_template('index.html', final_text=final_text)

@app.route('/download_doc', methods=['POST'])
def download_doc():
    global final_text

    # Erstelle ein neues Word-Dokument
    doc = Document()
    doc.add_heading('Vollständige Erinnerung', level=1)
    doc.add_paragraph(final_text)

    # Speichere das Dokument in einem BytesIO-Objekt
    doc_io = BytesIO()
    doc.save(doc_io)
    doc_io.seek(0)

    return send_file(doc_io, as_attachment=True, download_name='Erinnerung.doc', mimetype='application/vnd.openxmlformats-officedocument.wordprocessingml.document')

if __name__ == '__main__':
    app.run(debug=True)
